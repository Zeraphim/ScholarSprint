import unittest
import zipfile
from io import BytesIO
from unittest.mock import patch

import exporters
from exporters import (
    BULLETS,
    DOCX,
    HEADING,
    LATEX,
    MARKDOWN,
    NUMBERS,
    PARAGRAPH,
    PDF,
    RULE,
    TEXT,
    build_export,
    export_file_name,
    parse_inline,
    parse_markdown,
)

SAMPLE = """## Paper

study.pdf

## Executive Brief

A **bold** claim with *emphasis*, `inline code`, and a [source](https://example.com/a).

### Methods and Evidence

- First finding
- Second finding

1. Step one
2. Step two

---

### Limitations and Risks

Sample size was small (n < 30) & funding was 100% industry.
"""


class InlineParsingTests(unittest.TestCase):
    def test_formatting_markers(self):
        spans = parse_inline("plain **bold** *italic* `code` ***both***")
        self.assertEqual(
            [(span.text, span.bold, span.italic, span.code) for span in spans],
            [
                ("plain ", False, False, False),
                ("bold", True, False, False),
                (" ", False, False, False),
                ("italic", False, True, False),
                (" ", False, False, False),
                ("code", False, False, True),
                (" ", False, False, False),
                ("both", True, True, False),
            ],
        )

    def test_links_keep_text_and_url(self):
        (span,) = parse_inline("[source](https://example.com/a)")
        self.assertEqual((span.text, span.link), ("source", "https://example.com/a"))

    def test_unmatched_markers_are_literal(self):
        (span,) = parse_inline("2 * 3 = 6")
        self.assertEqual(span.text, "2 * 3 = 6")


class BlockParsingTests(unittest.TestCase):
    def setUp(self):
        self.blocks = parse_markdown(SAMPLE)

    def test_block_kinds_in_order(self):
        self.assertEqual(
            [block.kind for block in self.blocks],
            [
                HEADING,
                PARAGRAPH,
                HEADING,
                PARAGRAPH,
                HEADING,
                BULLETS,
                NUMBERS,
                RULE,
                HEADING,
                PARAGRAPH,
            ],
        )

    def test_heading_levels_shift_to_start_at_one(self):
        levels = [block.level for block in self.blocks if block.kind == HEADING]
        self.assertEqual(levels, [1, 1, 2, 2])

    def test_list_items_are_parsed(self):
        bullets = next(block for block in self.blocks if block.kind == BULLETS)
        numbers = next(block for block in self.blocks if block.kind == NUMBERS)
        self.assertEqual(
            ["".join(span.text for span in item) for item in bullets.items],
            ["First finding", "Second finding"],
        )
        self.assertEqual(
            ["".join(span.text for span in item) for item in numbers.items],
            ["Step one", "Step two"],
        )

    def test_fenced_code_is_kept_verbatim(self):
        blocks = parse_markdown("```\nx = 1\ny = 2\n```")
        self.assertEqual([block.kind for block in blocks], ["code"])
        self.assertEqual(blocks[0].text, "x = 1\ny = 2")

    def test_empty_input_yields_no_blocks(self):
        self.assertEqual(parse_markdown(""), [])


class MarkdownExportTests(unittest.TestCase):
    def test_title_is_prepended_and_body_preserved(self):
        output = build_export(SAMPLE, "study", MARKDOWN).decode("utf-8")
        self.assertTrue(output.startswith("# study\n\n"))
        self.assertIn("## Executive Brief", output)
        self.assertIn("[source](https://example.com/a)", output)


class TextExportTests(unittest.TestCase):
    def test_markdown_syntax_is_stripped(self):
        output = build_export(SAMPLE, "study", TEXT).decode("utf-8")
        self.assertIn("study\n=====", output)
        self.assertIn("A bold claim with emphasis", output)
        self.assertIn("- First finding", output)
        self.assertIn("1. Step one", output)
        self.assertNotIn("**", output)
        self.assertNotIn("###", output)


class LatexExportTests(unittest.TestCase):
    def setUp(self):
        self.output = build_export(SAMPLE, "study.pdf", LATEX).decode("utf-8")

    def test_document_skeleton(self):
        self.assertIn(r"\documentclass[11pt,a4paper]{article}", self.output)
        self.assertIn(r"\title{study.pdf}", self.output)
        self.assertTrue(self.output.rstrip().endswith(r"\end{document}"))

    def test_structure_and_formatting_commands(self):
        self.assertIn(r"\section{Paper}", self.output)
        self.assertIn(r"\subsection{Methods and Evidence}", self.output)
        self.assertIn(r"\textbf{bold}", self.output)
        self.assertIn(r"\textit{emphasis}", self.output)
        self.assertIn(r"\texttt{inline code}", self.output)
        self.assertIn(r"\href{https://example.com/a}{source}", self.output)
        self.assertIn(r"\begin{itemize}", self.output)
        self.assertIn(r"\begin{enumerate}", self.output)

    def test_special_characters_are_escaped(self):
        self.assertIn(r"100\% industry", self.output)
        self.assertIn(r"30) \& funding", self.output)

    def test_list_items_are_not_separated_by_blank_lines(self):
        self.assertIn("\\begin{itemize}\n  \\item First finding\n  \\item Second finding\n\\end{itemize}", self.output)

    def test_unicode_symbols_become_latex_commands(self):
        source = "Values \u2265 5, \u03b1 = 0.05, \u03c3 \u00b1 1.2, 25 \u00b0C, em\u2014dash, \u201cquoted\u201d, cost \u2248 3"
        output = build_export(source, "s", LATEX).decode("utf-8")
        for expected in (r"$\geq$", r"$\alpha$", r"$\pm$", r"\textdegree{}", "em---dash", "``quoted''", r"$\approx$"):
            self.assertIn(expected, output)
        self.assertNotIn("\u2265", output)
        self.assertNotIn("\u03b1", output)


class PdfExportTests(unittest.TestCase):
    def test_pdf_bytes_are_well_formed(self):
        output = build_export(SAMPLE, "study", PDF)
        self.assertTrue(output.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", output[-2048:])
        self.assertGreater(len(output), 1000)

    def test_unicode_punctuation_does_not_raise(self):
        output = build_export("## Heading\n\nEm\u2014dash \u201cquotes\u201d \u2265 90% \u2192 ok", "s", PDF)
        self.assertTrue(output.startswith(b"%PDF-"))

    def test_latin1_sanitizer(self):
        self.assertEqual(exporters.pdf_safe_text("a\u2014b\u2019c\u2026"), "a--b'c...")
        self.assertEqual(exporters.pdf_safe_text("\u03b1 \u2248 \u03c3"), "alpha ~ sigma")
        self.assertNotIn("?", exporters.pdf_safe_text("BLEU \u2265 28.4 \u2248 \u0394"))

    def test_extracted_pdf_text_matches_the_summary(self):
        from pypdf import PdfReader
        from io import BytesIO as _BytesIO

        reader = PdfReader(_BytesIO(build_export(SAMPLE, "study", PDF)))
        text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("study", text)
        self.assertIn("Executive Brief", text)
        self.assertIn("First finding", text)
        self.assertIn("Step two", text)


class DocxExportTests(unittest.TestCase):
    def setUp(self):
        self.output = build_export(SAMPLE, "study", DOCX)

    def test_docx_is_a_valid_package(self):
        self.assertTrue(self.output.startswith(b"PK"))
        with zipfile.ZipFile(BytesIO(self.output)) as archive:
            self.assertIn("word/document.xml", archive.namelist())

    def test_content_and_styles_round_trip(self):
        from docx import Document

        document = Document(BytesIO(self.output))
        texts = [paragraph.text for paragraph in document.paragraphs]
        styles = [paragraph.style.name for paragraph in document.paragraphs]

        self.assertIn("study", texts)
        self.assertIn("Executive Brief", texts)
        self.assertIn("First finding", texts)
        self.assertIn("List Bullet", styles)
        self.assertIn("List Number", styles)
        self.assertTrue(any(style.startswith("Heading") for style in styles))

        brief = next(p for p in document.paragraphs if p.text.startswith("A bold claim"))
        self.assertTrue(any(run.bold and run.text == "bold" for run in brief.runs))
        self.assertTrue(any(run.italic and run.text == "emphasis" for run in brief.runs))
        self.assertIn("source (https://example.com/a)", brief.text)


class FormatRegistryTests(unittest.TestCase):
    def test_every_format_is_available_after_sync(self):
        self.assertEqual(sorted(exporters.available_formats()), sorted(exporters.EXPORT_FORMATS))

    def test_file_names_use_the_stem_and_extension(self):
        self.assertEqual(export_file_name("study.pdf", MARKDOWN), "study.md")
        self.assertEqual(export_file_name("study.pdf", PDF), "study.pdf")
        self.assertEqual(export_file_name("a b/c.pdf", LATEX), "c.tex")
        self.assertEqual(export_file_name("", DOCX), "summary.docx")

    def test_unknown_format_is_rejected(self):
        with self.assertRaises(ValueError):
            build_export(SAMPLE, "study", "Excel (.xlsx)")

    def test_missing_dependency_reports_pip_name(self):
        with patch("importlib.util.find_spec", return_value=None):
            self.assertEqual(exporters.missing_dependency(PDF), "fpdf2")
            self.assertEqual(exporters.missing_dependency(DOCX), "python-docx")
        self.assertIsNone(exporters.missing_dependency(MARKDOWN))


if __name__ == "__main__":
    unittest.main()
